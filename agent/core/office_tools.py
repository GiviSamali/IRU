"""Bounded Office operations; no generated code, Office COM or cloud services."""
from __future__ import annotations

from copy import copy
from datetime import date, datetime
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
import threading
import zipfile

_LOCK = threading.RLock()
TOOLS = {"office_docx":"docx", "office_xlsx":"xlsx", "office_pptx":"pptx"}


class OfficeError(ValueError):
    pass


def require(condition, code):
    if not condition:
        raise OfficeError(code)


def fingerprint(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for data in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(data)
    return digest.hexdigest()


def text(value, maximum=12000):
    require(isinstance(value,str) and len(value) <= maximum, "invalid_text")
    return value


def page(items, selector):
    offset, limit = selector.get("offset",0), selector.get("limit",10)
    require(type(offset) is int and offset >= 0 and type(limit) is int and 1 <= limit <= 50, "invalid_page")
    result, size = [], 0
    for item in items[offset:offset + limit]:
        encoded = json.dumps(item,ensure_ascii=False,default=str)
        if result and size + len(encoded) > 10000:
            break
        require(len(encoded) <= 10000, "inspect_item_too_large_select_smaller_range")
        result.append(item)
        size += len(encoded)
    end = offset + len(result)
    return {"items":result, "offset":offset, "next_offset":end if end < len(items) else None,
            "total_items":len(items)}


def safe_edit_package(path, fmt):
    # Libraries cannot promise lossless round-tripping these parts in v1.
    forbidden = ("vbaproject", "activex/", "embeddings/", "_xmlsignatures/", "customxml/")
    if fmt == "xlsx":
        forbidden += ("drawings/", "charts/", "pivottable", "pivotcache", "slicer", "externallinks/",
                      "connections.xml", "querytables/", "richdata/", "threadedcomments/", "vml", "ctrlprops/")
    with zipfile.ZipFile(path) as archive:
        require(not any(any(token in name.lower() for token in forbidden) for name in archive.namelist()),
                "unsupported_edit_complex_package")
        if fmt == "xlsx":
            require(not any(b"extLst" in archive.read(name) for name in archive.namelist()
                if name.endswith(".xml")),"unsupported_edit_extension_data")
        if fmt == "pptx":
            import xml.etree.ElementTree as ET
            for name in archive.namelist():
                if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                    root=ET.fromstring(archive.read(name))
                    require(not any(node.tag.rsplit("}",1)[-1] in {"timing","transition"} for node in root.iter()),
                            "unsupported_edit_animated_presentation")
        if fmt == "docx":
            import xml.etree.ElementTree as ET
            root=ET.fromstring(archive.read("word/document.xml"))
            require(not any(node.tag.rsplit("}",1)[-1] in {"ins","del","sdt","altChunk"} for node in root.iter()),
                    "unsupported_edit_tracked_or_embedded_content")


def _docx_elements(doc):
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    elements = []
    for child in doc.element.body:
        if child.tag.endswith("}p"):
            elements.append(Paragraph(child,doc))
        elif child.tag.endswith("}tbl"):
            elements.append(Table(child,doc))
    return elements


def _docx_add(doc, elements):
    require(isinstance(elements,list) and len(elements) <= 200, "invalid_elements")
    for item in elements:
        kind = item.get("type","paragraph")
        if kind == "table":
            rows = item.get("rows",[])
            require(rows and len(rows) <= 200 and all(isinstance(r,list) and 1 <= len(r) <= 20 for r in rows), "invalid_table")
            width = len(rows[0])
            require(all(len(r)==width for r in rows), "ragged_table")
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for i,row in enumerate(rows):
                for j,value in enumerate(row):
                    table.cell(i,j).text = text(str(value),2000)
        elif kind == "page_break":
            doc.add_page_break()
        else:
            require(kind in {"paragraph","heading","bullet","numbered"}, "unsupported_docx_element")
            level = item.get("level",1)
            require(type(level) is int and 0 <= level <= 9, "invalid_heading_level")
            paragraph = doc.add_heading(level=level) if kind == "heading" else doc.add_paragraph(
                style={"bullet":"List Bullet","numbered":"List Number"}.get(kind,"Normal"))
            runs = item.get("runs",[{"text":item.get("text","")}])
            require(isinstance(runs,list) and len(runs) <= 100, "invalid_runs")
            for run in runs:
                target = paragraph.add_run(text(run.get("text","")))
                target.bold, target.italic = bool(run.get("bold",False)), bool(run.get("italic",False))


def docx_operation(path, action, data, changes, selector):
    from docx import Document
    from docx.shared import Pt
    from docx.text.paragraph import Paragraph
    doc = Document() if action == "create" else Document(str(path))
    if action == "create":
        style = data.get("style",{})
        doc.styles["Normal"].font.name = text(style.get("font","Calibri"),80)
        size = style.get("font_size",11)
        require(isinstance(size,(int,float)) and 8 <= size <= 36, "invalid_font_size")
        doc.styles["Normal"].font.size = Pt(size)
        _docx_add(doc,data.get("elements",[]))
    if action == "edit":
        original = _docx_elements(doc)
        untouched = {i:e._element.xml for i,e in enumerate(original)}
        cells={(i,r,c):cell._tc.xml for i,element in enumerate(original) if not isinstance(element,Paragraph)
               for r,row in enumerate(element.rows) for c,cell in enumerate(row.cells)}
        for change in changes:
            op = change.get("op")
            if op == "append":
                _docx_add(doc,change.get("elements",[]))
                continue
            element_id = change.get("element_id","")
            require(re.fullmatch(r"body:\d+",element_id) is not None, "invalid_element_id")
            index = int(element_id.split(":")[1])
            require(index < len(original), "element_not_found")
            element = original[index]
            untouched.pop(index,None)
            if op == "set_text":
                require(isinstance(element,Paragraph), "element_not_paragraph")
                require(not element._p.xpath(".//w:drawing|.//w:object|.//w:fldChar|.//w:hyperlink|.//w:fldSimple|.//w:bookmarkStart|.//w:footnoteReference|.//w:endnoteReference|.//w:commentReference"),
                        "unsupported_edit_complex_paragraph")
                # Only the selected paragraph's contents are replaced.
                element.text = text(change.get("text",""))
            elif op == "set_table_cell":
                require(not isinstance(element,Paragraph), "element_not_table")
                row,col = change.get("row"),change.get("column")
                require(type(row) is int and type(col) is int and 0 <= row < len(element.rows)
                        and 0 <= col < len(element.columns), "cell_out_of_range")
                cell=element.cell(row,col)
                require(not cell._tc.xpath("./w:tcPr/w:gridSpan|./w:tcPr/w:vMerge"),"unsupported_merged_table_edit")
                require(not cell.tables and not cell._tc.xpath(".//w:drawing|.//w:object|.//w:hyperlink|.//w:fldChar|.//w:fldSimple"),
                        "unsupported_edit_complex_table_cell")
                cells.pop((index,row,col),None)
                cell.text = text(change.get("text",""),2000)
            else:
                raise OfficeError("unsupported_docx_change")
        doc._iru_guard = {"docx":untouched,"docx_cells":cells}
    items = []
    selected = selector.get("element_id")
    for i,element in enumerate(_docx_elements(doc)):
        eid = "body:" + str(i)
        if selected and eid != selected:
            continue
        if isinstance(element,Paragraph):
            start = selector.get("text_offset",0)
            require(type(start) is int and start >= 0, "invalid_text_offset")
            items.append({"element_id":eid,"type":"paragraph","style":element.style.name,
                "text":element.text[start:start+1600],"total_chars":len(element.text),
                "next_text_offset":start+1600 if start+1600 < len(element.text) else None})
        else:
            row = selector.get("row_offset",0)
            col = selector.get("column_offset",0)
            start = selector.get("text_offset",0)
            require(type(row) is int and type(col) is int and type(start) is int and row >= 0 and col >= 0 and start >= 0, "invalid_table_offset")
            block=[[c.text for c in r.cells[col:col+6]] for r in element.rows[row:row+6]]
            items.append({"element_id":eid,"type":"table","rows":len(element.rows),"columns":len(element.columns),
                "row_offset":row,"column_offset":col,
                "cells":[[value[start:start+160] for value in r] for r in block],
                "cell_chars":[[len(value) for value in r] for r in block],"text_offset":start,
                "next_row_offset":row+6 if row+6 < len(element.rows) else None,
                "next_column_offset":col+6 if col+6 < len(element.columns) else None})
    require(not selected or bool(items), "element_not_found")
    return doc,{"paragraphs":len(doc.paragraphs),"tables":len(doc.tables)},page(items,selector)


def _sheet_name(value):
    name = text(value,31)
    require(bool(name) and not re.search(r"[\\/*?:\[\]]",name), "invalid_sheet_name")
    return name


def _cell_value(value):
    if isinstance(value,dict):
        require(set(value) in ({"date"},{"datetime"}),"invalid_date_value")
        parsed = date.fromisoformat(value["date"]) if "date" in value else datetime.fromisoformat(value["datetime"])
        require(not isinstance(parsed,datetime) or parsed.tzinfo is None,"unsupported_timezone_date")
        return parsed
    require(value is None or isinstance(value,(str,int,float,bool)), "invalid_cell_value")
    if isinstance(value,str):
        require(len(value) <= 12000, "cell_text_too_long")
        if value.startswith("="):
            require(not re.search(r"\[|\]|\||\b(?:WEBSERVICE|HYPERLINK|RTD|CALL|REGISTER|EXEC)\s*\(",value,re.I),
                    "unsupported_external_formula")
    return value


def _cell_style(cell, style):
    from openpyxl.styles import PatternFill
    require(set(style) <= {"bold","italic","fill","color","number_format","wrap"}, "unsupported_cell_style")
    font = copy(cell.font)
    if "bold" in style: font.bold = bool(style["bold"])
    if "italic" in style: font.italic = bool(style["italic"])
    if "color" in style:
        require(re.fullmatch(r"[0-9A-Fa-f]{6}",str(style["color"])) is not None,"invalid_color")
        font.color = style["color"]
    cell.font = font
    if "fill" in style:
        require(re.fullmatch(r"[0-9A-Fa-f]{6}",str(style["fill"])) is not None,"invalid_color")
        cell.fill = PatternFill("solid",fgColor=style["fill"])
    if "number_format" in style: cell.number_format = text(style["number_format"],100)
    if "wrap" in style:
        alignment = copy(cell.alignment)
        alignment.wrap_text = bool(style["wrap"])
        cell.alignment = alignment


def _write_rows(sheet, rows, start="A1", style=None):
    from openpyxl.utils.cell import coordinate_to_tuple
    require(isinstance(rows,list) and sum(len(r) for r in rows if isinstance(r,list)) <= 10000, "range_too_large")
    first_row,first_col = coordinate_to_tuple(start)
    require(first_row >= 1 and first_col >= 1, "invalid_cell_reference")
    for i,row in enumerate(rows):
        require(isinstance(row,list), "invalid_rows")
        for j,value in enumerate(row):
            require(first_row+i <= 1048576 and first_col+j <= 16384,"range_out_of_bounds")
            cell = sheet.cell(first_row+i,first_col+j)
            cell.value = _cell_value(value)
            if style: _cell_style(cell,style)


def _sheet_setup(sheet, spec):
    from openpyxl.utils.cell import column_index_from_string
    for column,width in spec.get("column_widths",{}).items():
        require(1 <= column_index_from_string(column) <= 16384 and isinstance(width,(int,float)) and 1 <= width <= 100,"invalid_column_width")
        sheet.column_dimensions[column].width = width
    if "freeze_panes" in spec: sheet.freeze_panes = spec["freeze_panes"]
    if spec.get("autofilter"): sheet.auto_filter.ref = sheet.dimensions
    if spec.get("header"):
        for cell in sheet[1]: _cell_style(cell,{"bold":True,"fill":"203864","color":"FFFFFF","wrap":True})
    for formatting in spec.get("formats",[]):
        from openpyxl.utils.cell import range_boundaries
        a,b,c,d=range_boundaries(formatting["range"])
        require(all(type(n) is int for n in (a,b,c,d)) and 1 <= a <= c <= 16384
            and 1 <= b <= d <= 1048576 and (c-a+1)*(d-b+1) <= 10000,"range_too_large")
        for row in sheet.iter_rows(min_row=b,max_row=d,min_col=a,max_col=c):
            for cell in row: _cell_style(cell,formatting.get("style",{}))


def xlsx_operation(path, action, data, changes, selector):
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils.cell import range_boundaries
    book = Workbook() if action == "create" else load_workbook(path,data_only=False,keep_links=True,rich_text=True)
    if action == "create":
        specs = data.get("sheets",[])
        require(isinstance(specs,list) and 1 <= len(specs) <= 30,"invalid_sheets")
        book.remove(book.active)
        for spec in specs:
            name = _sheet_name(spec["name"])
            require(name.casefold() not in {s.casefold() for s in book.sheetnames},"duplicate_sheet")
            sheet = book.create_sheet(name)
            _write_rows(sheet,spec.get("rows",[]))
            _sheet_setup(sheet,spec)
    if action == "edit":
        old_cells={(s.title,r,c):(cell.value,tuple(cell._style) if cell.has_style else None)
                   for s in book for (r,c),cell in s._cells.items()}
        old_sheets=list(book.sheetnames)
        touched=set()
        for change in changes:
            op = change.get("op")
            if op == "add_sheet":
                name = _sheet_name(change["name"])
                require(name.casefold() not in {s.casefold() for s in book.sheetnames},"duplicate_sheet")
                sheet = book.create_sheet(name)
                _write_rows(sheet,change.get("rows",[]))
                _sheet_setup(sheet,change)
                continue
            require(change.get("sheet") in book.sheetnames,"sheet_not_found")
            sheet = book[change["sheet"]]
            if op in {"write_range","set_cell","append_rows"}:
                require(op == "append_rows" or bool(change.get("cell") if op == "set_cell" else change.get("start")),"edit_target_required")
                rows = [[change.get("value")]] if op == "set_cell" else change.get("rows",[])
                start = "A"+str(sheet.max_row+1) if op == "append_rows" else change.get("start",change.get("cell","A1"))
                from openpyxl.utils.cell import coordinate_to_tuple
                r,c=coordinate_to_tuple(start)
                touched.update((sheet.title,r+i,c+j) for i,row in enumerate(rows) for j in range(len(row)))
                _write_rows(sheet,rows,start,change.get("style"))
            elif op == "format_range":
                bounds = range_boundaries(change["range"])
                require(all(type(n) is int for n in bounds),"invalid_range")
                a,b,c,d = bounds
                require(1 <= a <= c <= 16384 and 1 <= b <= d <= 1048576 and (c-a+1)*(d-b+1) <= 10000,"range_too_large")
                for row in sheet.iter_rows(min_row=b,max_row=d,min_col=a,max_col=c):
                    for cell in row:
                        touched.add((sheet.title,cell.row,cell.column))
                        _cell_style(cell,change.get("style",{}))
            elif op == "format_sheet":
                require(not change.get("header") and not change.get("formats"),"use_format_range_for_cell_styles")
                _sheet_setup(sheet,change)
            else:
                raise OfficeError("unsupported_xlsx_change")
        book._iru_guard={"xlsx":{key:value for key,value in old_cells.items() if key not in touched},"sheets":old_sheets}
    sheet_info=[{"name":s.title,"rows":s.max_row,"columns":s.max_column} for s in book]
    summary = {"sheets":len(book.worksheets),"rows":sum(s.max_row for s in book),"formula_calculation":"not_performed"}
    if selector.get("range"):
        require(selector.get("sheet") in book.sheetnames,"sheet_not_found")
        a,b,c,d = range_boundaries(selector["range"])
        require(all(type(n) is int for n in (a,b,c,d)) and 1 <= a <= c <= 16384
                and 1 <= b <= d <= 1048576 and (c-a+1)*(d-b+1) <= 100,"inspect_range_too_large")
        sheet = book[selector["sheet"]]
        cells=[]
        offset=selector.get("text_offset",0)
        require(type(offset) is int and offset >= 0,"invalid_text_offset")
        for row in sheet.iter_rows(min_row=b,max_row=d,min_col=a,max_col=c):
            for cell in row:
                value=cell.value.isoformat() if isinstance(cell.value,(date,datetime)) else cell.value
                if cell.data_type == "f" and isinstance(getattr(value,"text",None),str): value=value.text
                if value is not None and not isinstance(value,(str,int,float,bool)): value=str(value)
                item={"element_id":sheet.title+"!"+cell.coordinate,"value":value,"formula":cell.data_type=="f","number_format":cell.number_format}
                if isinstance(value,str):
                    item.update(value=value[offset:offset+1600],total_chars=len(value),
                        next_text_offset=offset+1600 if offset+1600 < len(value) else None)
                cells.append(item)
        return book,summary,page(cells,selector)
    return book,summary,page(sheet_info,selector)


def _ppt_text(shape, paragraphs, size=20, color="243247"):
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    require(isinstance(paragraphs,list) and len(paragraphs) <= 30,"invalid_paragraphs")
    require(isinstance(size,(int,float)) and 8 <= size <= 48,"invalid_font_size")
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for i,item in enumerate(paragraphs):
        item = {"text":item} if isinstance(item,str) else item
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = text(item.get("text",""),2000)
        p.level = item.get("level",0)
        require(type(p.level) is int and 0 <= p.level <= 5,"invalid_bullet_level")
        p.font.name = "Calibri"
        p.font.size = Pt(size)
        p.font.bold = bool(item.get("bold",False))
        p.font.italic = bool(item.get("italic",False))
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(12)
        if item.get("bullet") or p.level:
            from pptx.oxml.xmlchemy import OxmlElement
            node=OxmlElement("a:buChar")
            node.set("char","•")
            p._p.get_or_add_pPr().append(node)


def _ppt_add_content(slide, spec, width=13.333, height=7.5):
    from pptx.util import Inches
    sx,sy=width/13.333,height/7.5
    require(not (spec.get("table") and spec.get("images")),"unsupported_slide_table_image_combination")
    if "title" in spec:
        shape = slide.shapes.add_textbox(Inches(.6*sx),Inches(.35*sy),Inches(12*sx),Inches(.8*sy))
        shape.name = "iru:title"
        _ppt_text(shape,[{"text":spec["title"],"bold":True}],30,"203864")
    if spec.get("paragraphs") or spec.get("text"):
        paragraphs=spec.get("paragraphs",[spec.get("text","")])
        length=sum(len(p if isinstance(p,str) else p.get("text","")) for p in paragraphs)
        require(length <= (250 if spec.get("table") else 800 if spec.get("images") else 1600),"slide_text_too_large_split_slides")
        body_width=5.5 if spec.get("images") else 11.8
        body_height=1 if spec.get("table") else 5.4
        shape = slide.shapes.add_textbox(Inches(.75*sx),Inches(1.5*sy),Inches(body_width*sx),Inches(body_height*sy))
        _ppt_text(shape,paragraphs,16 if spec.get("table") or length>900 else 20)
    if "table" in spec:
        rows = spec["table"]
        require(rows and len(rows) <= 12 and 1 <= len(rows[0]) <= 8 and all(len(r)==len(rows[0]) for r in rows),"invalid_slide_table")
        require(all(len(str(value)) <= (40 if len(rows)>8 else 80) for row in rows for value in row),"slide_table_text_too_large")
        y,h=(2.8,3.8) if spec.get("paragraphs") or spec.get("text") else (1.6,4.8)
        shape = slide.shapes.add_table(len(rows),len(rows[0]),Inches(.75*sx),Inches(y*sy),Inches(11.8*sx),Inches(h*sy))
        for i,row in enumerate(rows):
            for j,value in enumerate(row):
                _ppt_text(shape.table.cell(i,j),[str(value)],12)
    images=spec.get("images",[])
    require(isinstance(images,list) and len(images) <= 8,"invalid_images")
    occupied=[]
    for image in images:
        path = Path(image["path"])
        require(path.is_absolute() and path.is_file() and not path.is_symlink()
                and path.suffix.lower() in {".png",".jpg",".jpeg"},"image_not_supported_or_missing")
        x,y,w,h = (image.get(k,v) for k,v in (("x",7*sx if spec.get("text") or spec.get("paragraphs") else .75*sx),
            ("y",1.6*sy),("width",5.5*sx),("height",4.8*sy)))
        require(all(isinstance(n,(int,float)) for n in (x,y,w,h)) and x >= 0 and y >= 0
                and w > 0 and h > 0 and x+w <= width and y+h <= height,"image_out_of_bounds")
        require(not any(x < a+c and a < x+w and y < b+d and b < y+h for a,b,c,d in occupied),"overlapping_images_unsupported")
        occupied.append((x,y,w,h))
        slide.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    if "notes" in spec:
        frame = slide.notes_slide.notes_text_frame
        require(frame is not None,"notes_not_supported")
        notes=text(spec["notes"])
        if frame.text.strip(): frame.add_paragraph().text=notes
        else: frame.text=notes


def pptx_operation(path, action, data, changes, selector):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation() if action == "create" else Presentation(str(path))
    def add_slide(spec):
        require(len(deck.slide_layouts)>0,"slide_layout_missing")
        slide=deck.slides.add_slide(deck.slide_layouts[min(6,len(deck.slide_layouts)-1)])
        # Only a newly created slide's default placeholders are removed.
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        _ppt_add_content(slide,spec,deck.slide_width.inches,deck.slide_height.inches)
    if action == "create":
        deck.slide_width,deck.slide_height = Inches(13.333),Inches(7.5)
        specs = data.get("slides",[])
        require(isinstance(specs,list) and 1 <= len(specs) <= 100,"invalid_slides")
        for spec in specs: add_slide(spec)
    if action == "edit":
        untouched={(s.slide_id,shape.shape_id):shape._element.xml for s in deck.slides for shape in s.shapes}
        for change in changes:
            op = change.get("op")
            if op == "add_slide":
                add_slide(change.get("slide",{}))
                continue
            slide = next((s for s in deck.slides if s.slide_id == change.get("slide_id")),None)
            require(slide is not None,"slide_not_found")
            if op == "set_text":
                shape = next((s for s in slide.shapes if str(s.shape_id)==str(change.get("shape_id"))),None)
                require(shape is not None and shape.has_text_frame,"text_shape_not_found")
                require(not shape._element.xpath(".//a:hlinkClick|.//a:hlinkMouseOver|.//a:fld"),"unsupported_edit_linked_text")
                untouched.pop((slide.slide_id,shape.shape_id),None)
                _ppt_text(shape,change.get("paragraphs",[change.get("text","")]),change.get("font_size",20))
            elif op == "add_content":
                _ppt_add_content(slide,change.get("content",{}),deck.slide_width.inches,deck.slide_height.inches)
            else:
                raise OfficeError("unsupported_pptx_change")
        deck._iru_guard={"pptx":untouched}
    items=[]
    for slide in deck.slides:
        if selector.get("slide_id") is not None and selector["slide_id"] != slide.slide_id: continue
        if selector.get("shape_id") is None:
            items.append({"slide_id":slide.slide_id,"type":"slide","shapes":len(slide.shapes),
                "notes":slide.notes_slide.notes_text_frame.text[:1000]
                    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None else ""})
        for shape in slide.shapes:
            if selector.get("shape_id") is not None and str(selector["shape_id"]) != str(shape.shape_id): continue
            item={"slide_id":slide.slide_id,"shape_id":shape.shape_id,"name":shape.name,
                  "type":"text" if shape.has_text_frame else "table" if shape.has_table else "other"}
            if shape.has_text_frame:
                offset=selector.get("text_offset",0)
                require(type(offset) is int and offset >= 0,"invalid_text_offset")
                item.update(text=shape.text[offset:offset+1600],total_chars=len(shape.text),
                    next_text_offset=offset+1600 if offset+1600 < len(shape.text) else None)
            if shape.has_table:
                r,c=selector.get("row_offset",0),selector.get("column_offset",0)
                require(type(r) is int and type(c) is int and r >= 0 and c >= 0,"invalid_table_offset")
                item.update(rows=len(shape.table.rows),columns=len(shape.table.columns),
                    row_offset=r,column_offset=c,cells=[[cell.text[:160] for cell in list(row.cells)[c:c+6]] for row in list(shape.table.rows)[r:r+6]],
                    next_row_offset=r+6 if r+6 < len(shape.table.rows) else None,
                    next_column_offset=c+6 if c+6 < len(shape.table.columns) else None)
            items.append(item)
    require(not (selector.get("slide_id") is not None or selector.get("shape_id") is not None) or bool(items),"element_not_found")
    return deck,{"slides":len(deck.slides)},page(items,selector)


def verify_preservation(document, reopened):
    guard=getattr(document,"_iru_guard",{})
    if "docx" in guard:
        elements=_docx_elements(reopened)
        require(all(i < len(elements) and elements[i]._element.xml==xml for i,xml in guard["docx"].items()),"unsupported_preservation_docx")
        require(all(elements[i].cell(r,c)._tc.xml==xml for (i,r,c),xml in guard["docx_cells"].items()),"unsupported_preservation_docx_table")
    if "xlsx" in guard:
        require(reopened.sheetnames[:len(guard["sheets"])]==guard["sheets"],"unsupported_preservation_sheets")
        for (sheet,r,c),(value,style) in guard["xlsx"].items():
            cell=reopened[sheet].cell(r,c)
            current=tuple(cell._style) if cell.has_style else None
            require(cell.value==value and current==style,"unsupported_preservation_xlsx")
    if "pptx" in guard:
        shapes={(s.slide_id,shape.shape_id):shape._element.xml for s in reopened.slides for shape in s.shapes}
        require(all(shapes.get(key)==xml for key,xml in guard["pptx"].items()),"unsupported_preservation_pptx")


def office(tool, action, path, data=None, changes=None, selector=None, expected_sha256=None, _source_task_id=None):
    """The same local executor is shared by all three registered agent actions."""
    tmp=None
    committed=False
    try:
        require(tool in TOOLS and action in {"create","inspect","edit"},"unsupported_office_action")
        target=Path(path)
        require(target.is_absolute() and not target.is_symlink(),"invalid_office_path")
        require(target.suffix.lower()=="."+TOOLS[tool],"unsupported_office_format")
        data,changes,selector=data or {},changes or [],selector or {}
        require(isinstance(data,dict) and isinstance(changes,list) and len(changes)<=100 and isinstance(selector,dict),"invalid_office_parameters")
        allowed={"docx":{"style","elements"},"xlsx":{"sheets"},"pptx":{"slides"}}[TOOLS[tool]]
        require(set(data) <= allowed,"unsupported_creation_parameter")
        if action == "create": require(not target.exists(),"file_exists_use_edit")
        else: require(target.is_file(),"file_not_found")
        before=fingerprint(target) if target.is_file() else None
        if action == "edit":
            require(isinstance(expected_sha256,str) and before==expected_sha256,"revision_required_or_stale")
            require(bool(changes),"empty_edit")
        handler={"office_docx":docx_operation,"office_xlsx":xlsx_operation,"office_pptx":pptx_operation}[tool]
        with _LOCK:
            working=target
            if action == "edit":
                import shutil
                fd,tmp=tempfile.mkstemp(prefix=".iru-office-",suffix=target.suffix,dir=target.parent)
                os.close(fd)
                shutil.copyfile(target,tmp)
                require(fingerprint(Path(tmp))==before,"file_changed_during_snapshot")
                safe_edit_package(Path(tmp),TOOLS[tool])
                working=Path(tmp)
            document,summary,inspection=handler(working,action,data,changes,selector)
            if action == "inspect":
                require(fingerprint(target)==before,"file_changed_during_inspect")
                return {"status":"success","action":action,"format":TOOLS[tool],"path":str(target),
                        "sha256":before,"size_bytes":target.stat().st_size,"summary":summary,**inspection}
            target.parent.mkdir(parents=True,exist_ok=True)
            if tmp is None:
                fd,tmp=tempfile.mkstemp(prefix=".iru-office-",suffix=target.suffix,dir=target.parent)
                os.close(fd)
            document.save(tmp)
            with open(tmp,"r+b") as stream: os.fsync(stream.fileno())
            # Reopen with the actual format library before publishing any bytes.
            reopened,_,_=handler(Path(tmp),"inspect",{},[],{})
            verify_preservation(document,reopened)
            digest=fingerprint(Path(tmp))
            if action == "edit":
                require(not target.is_symlink() and fingerprint(target)==before,"file_changed_before_commit")
                import shutil
                shutil.copymode(target,tmp)
                os.replace(tmp,target)
                committed=True
            else:
                os.link(tmp,target)  # Exclusive creation, never overwrite a racing writer.
                committed=True
                os.unlink(tmp)
            require(fingerprint(target)==digest,"file_changed_after_commit")
            receipt={"type":"file","path":str(target),"verified":True,"size_bytes":target.stat().st_size,
                     "sha256":digest,"source_task_id":_source_task_id,"created_or_changed":before != digest,
                     "operation":action,
                     "artifact_role":"final","format":TOOLS[tool]}
            return {"status":"success","action":action,"format":TOOLS[tool],"path":str(target),
                    "verified":True,"artifact_role":"final","size_bytes":receipt["size_bytes"],
                    "sha256":digest,"source_task_id":_source_task_id,"summary":summary,
                    "files_verified":[str(target)],"artifact_receipts":[receipt]}
    except ImportError:
        return {"status":"error","error":"office_dependencies_missing","required_packages":["python-docx","openpyxl","python-pptx"],
                "next_action":"update_agent_environment_or_rebuild_with_office_dependencies"}
    except Exception as exc:
        code = str(exc) if isinstance(exc,OfficeError) else "permission_denied" if isinstance(exc,PermissionError) else (
            "file_exists_use_edit" if isinstance(exc,FileExistsError) else "file_not_found" if isinstance(exc,FileNotFoundError)
            else "invalid_office_parameters" if isinstance(exc,(TypeError,ValueError,KeyError,IndexError)) else "office_operation_failed")
        return {"status":"unknown" if committed else "error","error":code,"error_type":type(exc).__name__,"committed":committed}
    finally:
        if tmp and os.path.exists(tmp):
            try: os.unlink(tmp)
            except OSError: pass
