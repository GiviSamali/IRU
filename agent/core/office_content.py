"""Content snapshots for formatting-only Office edits."""
from collections import Counter
import hashlib
import re
import zipfile
from pathlib import Path


def _text(values):
    # Preserve word order inside each paragraph; allow moving whole text blocks.
    return sorted(Counter(" ".join(line.split()) for value in values
                          for line in value.splitlines() if line.strip()).items())


def content_snapshot(path):
    path = Path(path)
    if not path.is_file() or path.is_symlink():
        raise ValueError("preservation_requires_existing_regular_file")
    if path.suffix.lower() == ".pptx":
        from pptx import Presentation
        def texts(shapes):
            out = []
            for shape in shapes:
                if shape.has_text_frame:
                    out.append(shape.text)
                if shape.has_table:
                    out.append(repr([[" ".join(cell.text.split()) for cell in row.cells]
                                     for row in shape.table.rows]))
                if shape.shape_type == 6:
                    out.extend(texts(shape.shapes))
            return out
        content = []
        for slide in Presentation(path).slides:
            notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ""
            content.append((_text(texts(slide.shapes)), _text([notes])))
    elif path.suffix.lower() == ".docx":
        import xml.etree.ElementTree as ET
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(path) as archive:
            names = sorted(n for n in archive.namelist() if n == "word/document.xml"
                           or re.fullmatch(r"word/(header|footer|footnotes|endnotes)[0-9]*\.xml", n))
            content = []
            for name in names:
                root = ET.fromstring(archive.read(name))
                paragraphs = _text(["".join(t.text or "" for t in p.iter(ns+"t"))
                                    for p in root.iter(ns+"p")])
                tables = [[[" ".join("".join(t.text or "" for t in cell.iter(ns+"t")).split())
                            for cell in row.findall(ns+"tc")]
                           for row in table.findall(ns+"tr")] for table in root.iter(ns+"tbl")]
                content.append((name, paragraphs, tables))
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, data_only=False, keep_links=True)
        try:
            content = [(s.title, [(cell.coordinate, type(cell.value).__name__, str(cell.value))
                        for _, cell in sorted(s._cells.items()) if cell.value is not None]) for s in book]
        finally:
            book.close()
    else:
        raise ValueError("preservation_unsupported_format")
    with zipfile.ZipFile(path) as archive:
        media = sorted(hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist()
                       if "/media/" in n and not n.endswith("/"))
        # Chart/embedded-object edits cannot be certified as formatting-only here.
        objects = sorted((n, hashlib.sha256(archive.read(n)).hexdigest()) for n in archive.namelist()
                         if ("/charts/" in n or "/embeddings/" in n) and not n.endswith("/"))
        import xml.etree.ElementTree as ET
        links = sorted((n, rel.get("Id"), rel.get("Target"))
                       for n in archive.namelist() if n.endswith(".rels")
                       for rel in ET.fromstring(archive.read(n)) if rel.get("TargetMode") == "External")
    return content, media, objects, links
